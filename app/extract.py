"""Turn uploaded files into content blocks the LLM can read.

- PDF with a text layer  -> plain text (cheap, fast)
- Scanned PDF            -> sent to Claude as a native PDF document block
- .docx                  -> plain text (paragraphs + tables)
- Images                 -> sent to Claude as image blocks (vision)
- .txt / .md             -> plain text
"""
import base64
import io

import pdfplumber
from docx import Document

IMAGE_TYPES = {
    "image/png": "image/png",
    "image/jpeg": "image/jpeg",
    "image/jpg": "image/jpeg",
    "image/webp": "image/webp",
    "image/gif": "image/gif",
}
MAX_TEXT_CHARS = 120_000


def _pdf_text(data: bytes) -> str:
    parts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
            # tables often hold the weekly schedule; keep them row-by-row
            for table in page.extract_tables() or []:
                for row in table:
                    parts.append(" | ".join(c or "" for c in row))
    return "\n".join(parts)


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def file_to_blocks(filename: str, content_type: str, data: bytes) -> list[dict]:
    name = (filename or "").lower()
    ctype = (content_type or "").lower()
    b64 = lambda: base64.standard_b64encode(data).decode()

    if name.endswith(".pdf") or ctype == "application/pdf":
        text = _pdf_text(data)
        if len(text.strip()) > 200:
            return [{"type": "text", "text": f"=== File: {filename} ===\n{text[:MAX_TEXT_CHARS]}"}]
        # scanned PDF: let Claude read the pages directly
        return [
            {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": b64()}},
            {"type": "text", "text": f"(The document above is {filename}, a scanned PDF.)"},
        ]

    if name.endswith(".docx"):
        return [{"type": "text", "text": f"=== File: {filename} ===\n{_docx_text(data)[:MAX_TEXT_CHARS]}"}]

    if ctype in IMAGE_TYPES or name.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif")):
        media = IMAGE_TYPES.get(ctype) or (
            "image/png" if name.endswith(".png")
            else "image/webp" if name.endswith(".webp")
            else "image/gif" if name.endswith(".gif")
            else "image/jpeg"
        )
        return [
            {"type": "image", "source": {"type": "base64", "media_type": media, "data": b64()}},
            {"type": "text", "text": f"(The image above is {filename}, a screenshot/photo of a syllabus.)"},
        ]

    # fallback: treat as text
    return [{"type": "text", "text": f"=== File: {filename} ===\n{data.decode('utf-8', 'ignore')[:MAX_TEXT_CHARS]}"}]
